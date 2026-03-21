from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

root = Path('/Users/revy/workspace_revy/securities_monolithic')
out_dir = root / 'output' / 'doc'
img_dir = root / 'output' / 'playwright'
out_path = out_dir / 'web-ui-portfolio.docx'

sections = [
    ('홈 화면', img_dir / 'portfolio-home-desktop.png'),
    ('시장 화면', img_dir / 'portfolio-chart-desktop.png'),
    ('계좌 화면', img_dir / 'portfolio-account-desktop.png'),
    ('환전 화면', img_dir / 'portfolio-exchange-desktop.png'),
    ('거래 화면', img_dir / 'portfolio-trade-desktop.png'),
    ('주문 화면', img_dir / 'portfolio-orders-desktop.png'),
    ('모바일 홈 화면', img_dir / 'portfolio-home-mobile.png'),
    ('모바일 거래 화면', img_dir / 'portfolio-trade-mobile.png'),
    ('모바일 주문 화면', img_dir / 'portfolio-orders-mobile.png'),
]

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Arial'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
style.font.size = Pt(10.5)

section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Web UI 개선 포트폴리오')
r.bold = True
r.font.size = Pt(22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('securities_monolithic / web-ui')
r.font.size = Pt(12)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run('고객용 금융 포털 화면을 실제 서비스형 UI로 재정리한 작업 정리 문서').italic = True

doc.add_paragraph('')

doc.add_heading('1. 프로젝트 개요', level=1)
for line in [
    '프로젝트명: securities_monolithic 고객용 프론트엔드 web-ui',
    '프로젝트 성격: 금융 서비스형 고객 포털 UI/UX 개선',
    '목표: 데모 느낌의 화면을 실제 서비스형 UI로 정리하고, PC와 모바일에서 모두 안정적으로 동작하도록 개선',
]:
    doc.add_paragraph(line, style='List Bullet')

doc.add_heading('2. 담당한 작업', level=1)
for line in [
    '공통 헤더와 홈 화면을 단순한 고객 포털 구조로 재구성',
    '시장, 계좌, 환전, 거래, 주문 화면의 정보 구조와 문구를 사용자 중심으로 재정리',
    '모바일 주문 화면을 카드형 목록으로 전환해 가로 깨짐 문제 완화',
    '차트, 거래, 주문 흐름의 상태 표시와 오류 메시지를 읽기 쉽게 정리',
    '전역 입력창, 버튼, 셀렉트박스 스타일을 통일해 일관성 개선',
]:
    doc.add_paragraph(line, style='List Bullet')

doc.add_heading('3. 핵심 개선 포인트', level=1)
points = {
    '서비스형 톤으로 재정리': [
        '과한 소개성 문구와 장식 요소를 줄이고 필요한 메뉴만 빠르게 이동하는 구조로 변경',
        '카드, 버튼, 라운드, 그림자 밀도를 줄여 더 단정한 금융 서비스 UI로 정리',
    ],
    '업무 화면 중심 UX 개선': [
        '시장: 종목 검색, 현재가, 차트, 주문 이동 흐름을 단순화',
        '계좌: 계좌 현황 확인 후 개설, 입금, 출금, 이체 작업으로 자연스럽게 연결',
        '환전: 송금이 아닌 환전 계산 흐름 중심으로 정보 구조 재정리',
        '거래 및 주문: 오류 메시지와 상태 표현을 사람이 읽기 쉬운 방식으로 정리',
    ],
    '반응형 구조 재설계': [
        'PC는 다열 정보 구조 유지',
        '태블릿은 2열 중심으로 재배치',
        '모바일은 1열과 전체 폭 버튼 중심으로 재배치',
        '주문 화면은 모바일에서 테이블 대신 카드형 목록으로 변경',
    ],
}
for title, bullets in points.items():
    doc.add_paragraph(title, style='List Number')
    for bullet in bullets:
        doc.add_paragraph(bullet, style='List Bullet')

doc.add_heading('4. 사용 기술', level=1)
for line in ['React', 'Vite', 'CSS', 'Playwright 기반 화면 캡처']:
    doc.add_paragraph(line, style='List Bullet')

doc.add_page_break()
doc.add_heading('5. 결과 화면', level=1)

for title, image_path in sections:
    doc.add_heading(title, level=2)
    if image_path.exists():
        doc.add_picture(str(image_path), width=Inches(6.3))
        cap = doc.add_paragraph(title)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f'이미지 없음: {image_path.name}')


doc.add_page_break()
doc.add_heading('6. 포트폴리오 메시지', level=1)
for line in [
    '데모 스타일 화면을 실제 금융 서비스형 UI로 재정리할 수 있음',
    '화면 디자인뿐 아니라 정보 구조, 오류 메시지, 반응형, 업무 흐름까지 함께 개선 가능',
    '고객용 화면에서 PC와 모바일을 분리해 사용 맥락에 맞게 설계 가능',
]:
    doc.add_paragraph(line, style='List Bullet')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer)

doc.save(out_path)
print(out_path)
